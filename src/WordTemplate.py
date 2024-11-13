import sys
from datetime import datetime
from pathlib import Path
from typing import Dict, Optional

from docx import Document


class WordTemplate:
    """Lädt und manipuliert die Word-Dokument-Vorlage zur Erstellung von Berichten."""

    def __init__(self, logger, doc_path: Path):
        """
        Initialisiert das WordTemplate-Objekt.

        Parameter:
        - logger: Ein Logger-Objekt zur Protokollierung von Informationen und Fehlern.
        - doc_path: Ein Pfad zur Word-Vorlagendatei.
        """
        self.logger = logger
        self.doc_path = doc_path
        self.document = self.load_template()  # Lädt die Word-Vorlage beim Erstellen des Objekts

    def load_template(self) -> Optional[Document]:
        """
        Lädt die Word-Dokumentvorlage und gibt sie als Document-Objekt zurück.

        Rückgabewert:
        - Ein Document-Objekt, falls das Dokument erfolgreich geladen wurde.
        - None, falls das Dokument nicht gefunden oder ein Fehler aufgetreten ist.
        """
        try:
            return Document(self.doc_path)
        except Exception:
            # Protokolliert einen Fehler, wenn die Datei nicht gefunden wird oder nicht lesbar ist.
            self.logger.error(f"Datei nicht gefunden: {self.doc_path}")
            return None

    def replace_placeholders(self, cell, placeholder: str, value: str):
        """
        Ersetzt einen Platzhalter in einer Zelle des Dokuments durch einen Wert.

        Parameter:
        - cell: Die Tabellenzelle, in der die Platzhalter ersetzt werden sollen.
        - placeholder: Der Platzhaltertext, der im Zellinhalt gesucht wird.
        - value: Der Wert, der anstelle des Platzhalters eingefügt wird.
        """
        if placeholder in cell.text:
            # Ersetzt den Platzhalter nur, wenn er in der Zelle vorhanden ist.
            cell.text = cell.text.replace(placeholder, str(value))

    def replace_general_placeholders(self, cell, placeholders: Dict[str, str]):
        """
        Ersetzt mehrere Platzhalter in einer Zelle basierend auf einem Dictionary von Platzhaltern und Werten.

        Parameter:
        - cell: Die Tabellenzelle, in der die Platzhalter ersetzt werden sollen.
        - placeholders: Ein Dictionary mit Platzhaltern als Schlüssel und den zugehörigen Werten als Werte.
        """
        for placeholder, value in placeholders.items():
            # Iteriert durch alle Platzhalter und ersetzt sie in der Zelle.
            self.replace_placeholders(cell, placeholder, value)

    def save_document(self, output_folder: Path):
        """
        Speichert das Word-Dokument in einem angegebenen Ordner und fügt einen Zeitstempel zum Dateinamen hinzu.

        Parameter:
        - output_folder: Der Pfad zum Ordner, in dem das Dokument gespeichert werden soll.

        Rückgabewert:
        - Der vollständige Pfad der gespeicherten Datei.
        """
        output_folder.mkdir(parents=True, exist_ok=True)  # Erstellt den Ordner, falls er nicht existiert
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")  # Generiert einen Zeitstempel
        output_path = output_folder / f'Ausbildungsnachweis_{timestamp}.docx'  # Setzt den Dateinamen

        try:
            self.document.save(output_path)
            # Erfolgreiches Speichern wird protokolliert
            self.logger.info(f"Dokument gespeichert unter: {output_path}")
        except Exception as e:
            # Fehler beim Speichern wird protokolliert und das Programm beendet
            self.logger.error(f"Fehler beim Speichern: {str(e)}")
            sys.exit(1)  # Beendet das Programm bei einem Speichernfehler
        return output_path
