import sys
from datetime import datetime
from pathlib import Path
from typing import Dict, Optional

from docx import Document

"""
Lädt und manipuliert die Word-Dokument-Vorlage zur Erstellung von Nachweisen.
"""
class WordTemplate:

    """
    Initialisiert das WordTemplate-Objekt.
    """
    def __init__(self, doc_path: Path):
        self.document = Document(doc_path)

    """
    Ersetzt einen Platzhalter in einer Zelle des Dokuments durch einen Wert.
    """
    @staticmethod
    def replace_placeholders(cell, placeholder: str, value: str):
        if placeholder in cell.text:
            # Ersetzt den Platzhalter nur, wenn er in der Zelle vorhanden ist.
            cell.text = cell.text.replace(placeholder, str(value))

    """
    Ersetzt mehrere Platzhalter in einer Zelle basierend auf einem Dictionary von Platzhaltern und Werten.
    """
    def replace_general_placeholders(self, cell, placeholders: Dict[str, str]):
        for placeholder, value in placeholders.items():
            # Iteriert durch alle Platzhalter und ersetzt sie in der Zelle.
            self.replace_placeholders(cell, placeholder, value)

    """
    Speichert das Word-Dokument in einem angegebenen Ordner und fügt einen Zeitstempel zum Dateinamen hinzu
    """
    def save_document(self, output_folder: Path):
        output_folder.mkdir(parents=True, exist_ok=True)  # Erstellt den Ordner, falls er nicht existiert
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")  # Generiert einen Zeitstempel
        output_path = output_folder / f'Ausbildungsnachweis_{timestamp}.docx'  # Setzt den Dateinamen

        try:
            self.document.save(output_path)
        except Exception as e:
            # Fehler beim Speichern
            pass
        return output_path
